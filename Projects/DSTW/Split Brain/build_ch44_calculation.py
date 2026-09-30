from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


OUTPUT = Path(__file__).with_name("Split-Brain-Chapter-4.4-Calculations.xlsx")

BLUE = "D9EAF7"
GREEN = "E2F0D9"
ORANGE = "FCE4D6"
GREY = "E7E6E6"
NAVY = "1F4E78"
WHITE = "FFFFFF"
THIN_GREY = Side(style="thin", color="B7B7B7")


def style_title(sheet, cell_range, text):
    sheet.merge_cells(cell_range)
    cell = sheet[cell_range.split(":")[0]]
    cell.value = text
    cell.font = Font(bold=True, color=WHITE, size=14)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="left")


def style_header(row):
    for cell in row:
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = Border(bottom=THIN_GREY)


def style_table(sheet, min_row, max_row, min_col, max_col):
    for row in sheet.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.border = Border(bottom=THIN_GREY)
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def configure_widths(sheet, widths):
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def build_workbook():
    workbook = Workbook()
    assumptions = workbook.active
    assumptions.title = "Assumptions"
    calc = workbook.create_sheet("4.4 Calculation")
    summary = workbook.create_sheet("4.4 Summary")
    trace = workbook.create_sheet("8h Trace")
    all_trace = workbook.create_sheet("All MTTR Trace")
    guide = workbook.create_sheet("Formula Guide")

    style_title(assumptions, "A1:D1", "DSTW Split Brain - Chapter 4.4 Calculation Workbook")
    assumptions["A3"] = "Editable input"
    assumptions["B3"] = "Value"
    assumptions["C3"] = "Unit"
    assumptions["D3"] = "Purpose"
    style_header(assumptions[3])
    inputs = [
        ("Core-router FIT", 10000, "FIT", "Failure rate for each core router"),
        ("Single-router FIT", 3500, "FIT", "Failure rate for the single router"),
        ("CCF beta factor", 0.02, "-", "Fraction of each core-router rate assigned to CCF"),
        ("Hours per year", 8760, "h/year", "Conversion constant"),
        ("Minutes per year", "=B7*60", "min/year", "Conversion constant"),
        ("Seconds per minute", 60, "sec/min", "Conversion constant"),
    ]
    for row_number, values in enumerate(inputs, start=4):
        for column, value in enumerate(values, start=1):
            assumptions.cell(row_number, column, value)
        assumptions.cell(row_number, 2).fill = PatternFill("solid", fgColor=BLUE if row_number < 8 else GREY)
    assumptions["A12"] = "Model assumptions"
    assumptions["A12"].font = Font(bold=True)
    assumptions["A13"] = "1oo2 core routers are hot redundant; one independent failure is tolerated."
    assumptions["A14"] = "All MTTR values are assumed equal and the combined unavailable-state duration is approximated by MTTR."
    assumptions["A15"] = "The single router is assumed independent of the core-router pair."
    assumptions["A16"] = "The low-unavailability approximation is intended for RAM sensitivity analysis, not final safety justification."
    configure_widths(assumptions, {"A": 28, "B": 16, "C": 14, "D": 62})
    assumptions.freeze_panes = "A4"

    style_title(calc, "A1:L1", "Chapter 4.4 - Combined failure calculation")
    calc["A3"] = "MTTR (h)"
    headers = [
        "Core CCF rate (FIT)",
        "Core independent rate (FIT)",
        "U core CCF",
        "U core independent",
        "U core total",
        "U single router",
        "U triple total",
        "Triple rate (h^-1)",
        "Triple rate (FIT)",
        "Downtime (min/y)",
        "Downtime (sec/y)",
    ]
    for column, header in enumerate(headers, start=2):
        calc.cell(3, column, header)
    style_header(calc[3])
    mttrs = [1, 4, 8, 12, 24]
    for row_number, mttr in enumerate(mttrs, start=4):
        calc.cell(row_number, 1, mttr)
        calc.cell(row_number, 2, "=Assumptions!$B$4*Assumptions!$B$6")
        calc.cell(row_number, 3, "=Assumptions!$B$4*(1-Assumptions!$B$6)")
        calc.cell(row_number, 4, f"=B{row_number}/1000000000*A{row_number}")
        calc.cell(row_number, 5, f"=2*(C{row_number}/1000000000*A{row_number})^2")
        calc.cell(row_number, 6, f"=D{row_number}+E{row_number}")
        calc.cell(row_number, 7, f"=Assumptions!$B$5/1000000000*A{row_number}")
        calc.cell(row_number, 8, f"=F{row_number}*G{row_number}")
        calc.cell(row_number, 9, f"=H{row_number}/A{row_number}")
        calc.cell(row_number, 10, f"=I{row_number}*1000000000")
        calc.cell(row_number, 11, f"=H{row_number}*Assumptions!$B$8")
        calc.cell(row_number, 12, f"=K{row_number}*Assumptions!$B$9")
    style_table(calc, 4, 8, 1, 12)
    for row in range(4, 9):
        calc.cell(row, 1).fill = PatternFill("solid", fgColor=BLUE)
        for column in range(2, 13):
            calc.cell(row, column).fill = PatternFill("solid", fgColor=GREEN)
    calc["A10"] = "Formula interpretation"
    calc["A10"].font = Font(bold=True)
    calc["A11"] = "U core total = core CCF unavailability + independent 1oo2 double-failure unavailability."
    calc["A12"] = "U triple total = U core total x U single router."
    calc["A13"] = "The single-router term is required because total network failure needs both the core-router pair and the single router unavailable."
    configure_widths(calc, {"A": 13, "B": 18, "C": 23, "D": 15, "E": 20, "F": 15, "G": 15, "H": 15, "I": 17, "J": 18, "K": 18, "L": 17})
    calc.freeze_panes = "A4"
    calc.auto_filter.ref = "A3:L8"
    calc.conditional_formatting.add("L4:L8", CellIsRule(operator="greaterThan", formula=["0.001"], fill=PatternFill("solid", fgColor=ORANGE)))

    summary_headers = [
        "FiT",
        "MTBF (h)",
        "MTBF (y)",
        "MTTR (h)",
        "Availability",
        "Unavailability",
        "Downtime (min/year)",
        "Downtime (sec/year)",
    ]
    style_title(summary, "A1:H1", "Chapter 4.4 Summary - Requested RAM fields")
    summary["A2"] = "1oo2 core-router pair including CCF, beta = 0.02"
    summary["A2"].font = Font(bold=True)
    for column, header in enumerate(summary_headers, start=1):
        summary.cell(3, column, header)
    style_header(summary[3])
    for row_number, source_row in enumerate(range(4, 9), start=4):
        formulas = [
            f"=Assumptions!$B$4*Assumptions!$B$6+2*(Assumptions!$B$4*(1-Assumptions!$B$6)/1000000000)^2*D{row_number}*1000000000",
            f"=1000000000/A{row_number}",
            f"=B{row_number}/Assumptions!$B$7",
            f"='4.4 Calculation'!A{source_row}",
            f"=1-F{row_number}",
            f"=A{row_number}/1000000000*D{row_number}",
            f"=F{row_number}*Assumptions!$B$8",
            f"=G{row_number}*Assumptions!$B$9",
        ]
        for column, formula in enumerate(formulas, start=1):
            summary.cell(row_number, column, formula)
    style_table(summary, 4, 8, 1, 8)

    summary["A11"] = "1oo2 core-router pair without CCF"
    summary["A11"].font = Font(bold=True)
    for column, header in enumerate(summary_headers, start=1):
        summary.cell(12, column, header)
    style_header(summary[12])
    for row_number, source_row in enumerate(range(4, 9), start=13):
        formulas = [
            f"=2*(Assumptions!$B$4/1000000000)^2*D{row_number}*1000000000",
            f"=1000000000/A{row_number}",
            f"=B{row_number}/Assumptions!$B$7",
            f"='4.4 Calculation'!A{source_row}",
            f"=1-F{row_number}",
            f"=A{row_number}/1000000000*D{row_number}",
            f"=F{row_number}*Assumptions!$B$8",
            f"=G{row_number}*Assumptions!$B$9",
        ]
        for column, formula in enumerate(formulas, start=1):
            summary.cell(row_number, column, formula)
    style_table(summary, 13, 17, 1, 8)
    configure_widths(summary, {"A": 16, "B": 18, "C": 16, "D": 13, "E": 18, "F": 20, "G": 22, "H": 22})
    summary.freeze_panes = "A4"

    chart = LineChart()
    chart.title = "Combined downtime versus MTTR"
    chart.y_axis.title = "Downtime (sec/year)"
    chart.x_axis.title = "MTTR (hours)"
    chart_data = Reference(calc, min_col=12, min_row=3, max_row=8)
    chart_categories = Reference(calc, min_col=1, min_row=4, max_row=8)
    chart.add_data(chart_data, titles_from_data=True)
    chart.set_categories(chart_categories)
    chart.height = 7
    chart.width = 13
    calc.add_chart(chart, "A16")

    style_title(trace, "A1:D1", "8-hour worked example")
    trace["A3"] = "Layer"
    trace["B3"] = "Calculation"
    trace["C3"] = "Excel formula / value"
    trace["D3"] = "Result"
    style_header(trace[3])
    trace_rows = [
        ("Inputs", "MTTR", "='4.4 Calculation'!A6", "8 h"),
        ("Layer 1", "Core CCF rate", "=Assumptions!B4*Assumptions!B6", "200 FIT"),
        ("Layer 1", "Core independent rate", "=Assumptions!B4*(1-Assumptions!B6)", "9800 FIT"),
        ("Layer 1", "U core CCF", "=C5/1000000000*C4", "1.60E-6"),
        ("Layer 1", "U core independent", "=2*(C6/1000000000*C4)^2", "1.229312E-8"),
        ("Layer 1", "U core total", "=C7+C8", "1.61229312E-6"),
        ("Layer 2", "Single-router rate", "=Assumptions!B5/1000000000", "3.5E-6 /h"),
        ("Layer 2", "U single router", "=C10*C4", "2.8E-5"),
        ("Layer 3", "U triple total", "=C9*C11", "4.514420736E-11"),
        ("Layer 3", "Combined failure rate", "=C12/C4", "5.643E-12 /h"),
        ("Layer 3", "Combined rate in FIT", "=C13*1000000000", "0.005643 FIT"),
        ("Layer 3", "Downtime sec/year", "=C12*Assumptions!B8*Assumptions!B9", "0.001424 sec/y"),
    ]
    for row_number, values in enumerate(trace_rows, start=4):
        for column, value in enumerate(values, start=1):
            trace.cell(row_number, column, value)
    style_table(trace, 4, 16, 1, 4)
    configure_widths(trace, {"A": 14, "B": 28, "C": 46, "D": 22})
    trace.freeze_panes = "A4"

    style_title(all_trace, "A1:I1", "Three-layer calculation for all MTTR values")
    all_trace_headers = [
        "MTTR (h)",
        "Layer 1: U core CCF",
        "Layer 1: U core independent",
        "Layer 1: U core total",
        "Layer 2: U single router",
        "Layer 3: U triple total",
        "Combined rate (FIT)",
        "Downtime (min/y)",
        "Downtime (sec/y)",
    ]
    for column, header in enumerate(all_trace_headers, start=1):
        all_trace.cell(3, column, header)
    style_header(all_trace[3])
    for row_number in range(4, 9):
        source_row = row_number
        formulas = [
            f"='4.4 Calculation'!A{source_row}",
            f"='4.4 Calculation'!D{source_row}",
            f"='4.4 Calculation'!E{source_row}",
            f"='4.4 Calculation'!F{source_row}",
            f"='4.4 Calculation'!G{source_row}",
            f"='4.4 Calculation'!H{source_row}",
            f"='4.4 Calculation'!J{source_row}",
            f"='4.4 Calculation'!K{source_row}",
            f"='4.4 Calculation'!L{source_row}",
        ]
        for column, formula in enumerate(formulas, start=1):
            all_trace.cell(row_number, column, formula)
    style_table(all_trace, 4, 8, 1, 9)
    all_trace["A10"] = "Each row follows the same sequence: core-router pair unavailability, single-router unavailability, then simultaneous overlap."
    all_trace["A10"].alignment = Alignment(wrap_text=True)
    configure_widths(all_trace, {"A": 13, "B": 20, "C": 27, "D": 20, "E": 21, "F": 20, "G": 20, "H": 20, "I": 20})
    all_trace.freeze_panes = "A4"

    style_title(guide, "A1:C1", "Formula guide")
    guide["A3"] = "Quantity"
    guide["B3"] = "Formula"
    guide["C3"] = "Meaning"
    style_header(guide[3])
    guide_rows = [
        ("Core CCF rate", "beta x lambda_CR", "CCF directly removes both core routers"),
        ("Core independent rate", "(1 - beta) x lambda_CR", "Independent portion of each core-router rate"),
        ("U core CCF", "lambda_CCF x MTTR", "Unavailability from common-cause core loss"),
        ("U core independent", "2 x (lambda_CR,ind x MTTR)^2", "Independent 1oo2 double failure; factor 2 for either first failure"),
        ("U core total", "U core CCF + U core independent", "Total core-router-pair unavailability"),
        ("U single router", "lambda_S x MTTR", "Unavailability of the single router"),
        ("U triple total", "U core total x U single router", "All required failure conditions overlap"),
        ("Combined failure rate", "U triple total / MTTR", "Equivalent rate using MTTR as combined exposure duration"),
        ("Downtime sec/year", "U triple total x 8760 x 3600", "Annual accumulated downtime in seconds"),
    ]
    for row_number, values in enumerate(guide_rows, start=4):
        for column, value in enumerate(values, start=1):
            guide.cell(row_number, column, value)
    style_table(guide, 4, 12, 1, 3)
    configure_widths(guide, {"A": 24, "B": 52, "C": 70})
    guide.freeze_panes = "A4"

    for sheet in workbook.worksheets:
        sheet.sheet_view.showGridLines = False
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.number_format = "0.000000000000"
    for row_number in list(range(4, 9)) + list(range(13, 18)):
        summary.cell(row_number, 1).number_format = "0.000"
        summary.cell(row_number, 2).number_format = "0.000"
        summary.cell(row_number, 3).number_format = "0.000"
        summary.cell(row_number, 4).number_format = "0"
        summary.cell(row_number, 5).number_format = "0.000000000"
        summary.cell(row_number, 6).number_format = "0.000000000000"
        summary.cell(row_number, 7).number_format = "0.000000"
        summary.cell(row_number, 8).number_format = "0.000000"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(OUTPUT)


if __name__ == "__main__":
    build_workbook()
    print(OUTPUT)